import os
import pandas as pd
from PyPDF2 import PdfReader
from docx import Document
from PIL import Image
import pytesseract
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# ---------------- Helper functions ----------------

def read_pdf(file_path):
    text = ""
    reader = PdfReader(file_path)
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + " "
    return text.strip()

def read_docx(file_path):
    doc = Document(file_path)
    return " ".join([p.text for p in doc.paragraphs])

def read_image(file_path):
    img = Image.open(file_path)
    return pytesseract.image_to_string(img)

def get_matched_skills(applicant_skills, employer_skills):
    a_skills = {s.strip().lower() for s in applicant_skills.split(",")}
    e_skills = {s.strip().lower() for s in employer_skills.split(",")}
    return ", ".join(sorted(a_skills & e_skills)) if a_skills & e_skills else "No Match"

# ---------------- Load resumes ----------------
resume_folder = "resumes"
data = []

for file in os.listdir(resume_folder):
    path = os.path.join(resume_folder, file)
    if file.lower().endswith(".pdf"):
        text = read_pdf(path)
    elif file.lower().endswith(".docx"):
        text = read_docx(path)
    elif file.lower().endswith((".jpg", ".jpeg", ".png")):
        text = read_image(path)
    else:
        continue
    data.append({"Name": os.path.splitext(file)[0], "Type": "Applicant", "Skills": text})

applicants = pd.DataFrame(data)

# ---------------- Example employers (can also load from CSV) ----------------
employers = pd.DataFrame([
    {"Name": "Company A", "Type": "Company", "Skills": "Python, Java, SQL, Communication"},
    {"Name": "Company B", "Type": "Company", "Skills": "PHP, HTML, CSS, JavaScript, Teamwork"}
])

# ---------------- TF-IDF + Cosine Similarity ----------------
all_docs = pd.concat([employers["Skills"], applicants["Skills"]])
vectorizer = TfidfVectorizer()
tfidf_matrix = vectorizer.fit_transform(all_docs)

employer_tfidf = tfidf_matrix[:len(employers)]
applicant_tfidf = tfidf_matrix[len(employers):]
similarity_matrix = cosine_similarity(applicant_tfidf, employer_tfidf)

# ---------------- Build merged dashboard ----------------
merged_dashboard = []

# Applicant view
for i, applicant in applicants.iterrows():
    similarities = similarity_matrix[i]
    top_indices = similarities.argsort()[::-1][:5]
    print(f"\nApplicant: {applicant['Name']} Top 3 Employers:")
    for rank, idx in enumerate(top_indices[:3], start=1):
        employer = employers.iloc[idx]
        percentage = similarities[idx] * 100
        matched_skills = get_matched_skills(applicant["Skills"], employer["Skills"])
        print(f"  Top {rank}: {employer['Name']} ({percentage:.2f}%) | Matched Skills: {matched_skills}")
        merged_dashboard.append({
            "View Type": "Applicant View",
            "Person": applicant["Name"],
            "Rank": rank,
            "Matched With": employer["Name"],
            "Match %": f"{percentage:.2f}",
            "Matched Skills": matched_skills
        })

# Employer view
for j, employer in employers.iterrows():
    similarities = similarity_matrix[:, j]
    top_indices = similarities.argsort()[::-1][:5]
    print(f"\nEmployer: {employer['Name']} Top 3 Applicants:")
    for rank, idx in enumerate(top_indices[:3], start=1):
        applicant = applicants.iloc[idx]
        percentage = similarities[idx] * 100
        matched_skills = get_matched_skills(applicant["Skills"], employer["Skills"])
        print(f"  Top {rank}: {applicant['Name']} ({percentage:.2f}%) | Matched Skills: {matched_skills}")
        merged_dashboard.append({
            "View Type": "Employer View",
            "Person": employer["Name"],
            "Rank": rank,
            "Matched With": applicant["Name"],
            "Match %": f"{percentage:.2f}",
            "Matched Skills": matched_skills
        })

# ---------------- Save to Excel with Top 1 highlight ----------------
merged_df = pd.DataFrame(merged_dashboard)
merged_df['Highlight'] = merged_df['Rank'].apply(lambda x: "Yes" if x == 1 else "")
excel_file = "merged_dashboard_resumes.xlsx"
merged_df.to_excel(excel_file, index=False)

# Highlight Top 1 in Excel
wb = load_workbook(excel_file)
ws = wb.active
fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
highlight_col = None
for i, cell in enumerate(ws[1], start=1):
    if cell.value == "Highlight":
        highlight_col = i
        break
if highlight_col:
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if row[highlight_col-1].value == "Yes":
            for cell in row:
                cell.fill = fill
wb.save(excel_file)

print(f"\nDashboard saved as {excel_file} with Top 1 highlighted.")
